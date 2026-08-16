"""B 系列(第二批外部审计修复)的回归测试。

覆盖:B1 buyer 匹配中间剥除碰撞(见 test_buyer_matcher.py)、
B3 PL 行净重自洽 + CI 读 PL 回写净重 + 跨单净重校验。
(B2 版本一致性检查见 test_check_version.py)
"""
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from tests.conftest import make_resolved_model, SAMPLE_CONFIG
from trade_pipeline.writers.pl_writer_lite import _compute_packing, PLWriterLite
from trade_pipeline.writers.ci_writer import CIWriter
from trade_pipeline.validation.cross_doc import check_ci_pl_net_weight
import pytest


def _item(**kw):
    from tests.conftest import make_model
    from dataclasses import replace
    base = make_model(with_weights=False).items[0]
    return replace(base, **kw)


# ── B3:行净重 = 每箱重 × 箱数 ────────────────────────────────────


def test_row_net_weight_is_carton_weight_times_cartons():
    # 100kg 分 3 箱(kg_per_carton_override=34 → ceil(100/34)=3):
    # 每箱重 round(100/3,2)=33.33,行净重必须是 33.33×3=99.99,
    # 不再是原来的真实净重 100.00(33.33×3≠100 报关算术不过)
    lines, summary = _compute_packing(
        [_item(quantity=1000, weight_kg=100.0, kg_per_carton_override=34.0)])
    line = lines[0]
    assert line.cartons == 3
    assert line.kg_per_carton == pytest.approx(33.33)
    assert line.net_weight_kg == pytest.approx(99.99)
    # 单据内部自洽恒等式
    assert round(line.kg_per_carton * line.cartons, 2) == line.net_weight_kg
    # 合计 = Σ行净重
    assert summary["total_net_weight"] == pytest.approx(99.99)


def test_total_net_is_sum_of_row_nets():
    items = [
        _item(quantity=1000, weight_kg=100.0, kg_per_carton_override=34.0),  # 99.99
        _item(quantity=1000, weight_kg=50.0, kg_per_carton_override=25.0),   # 50.00 整除
    ]
    lines, summary = _compute_packing(items)
    expect = round(sum(round(ln.kg_per_carton * ln.cartons, 2) for ln in lines), 2)
    assert summary["total_net_weight"] == pytest.approx(expect)
    assert summary["total_net_weight"] == pytest.approx(149.99)


def test_evenly_divisible_weight_unchanged():
    # 整除场景(50kg/2箱)不引入偏差:行净重仍等于真实净重
    lines, _ = _compute_packing(
        [_item(quantity=1000, weight_kg=50.0, kg_per_carton_override=25.0)])
    assert lines[0].kg_per_carton == 25.0
    assert lines[0].net_weight_kg == 50.0


def test_pl_excel_cells_self_consistent():
    """从 Excel 单元格直接断言:D列×C列==F列,ΣF==合计F(现有测试此前无 D 列断言)。"""
    model = make_resolved_model(with_weights=True)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        PLWriterLite(model, SAMPLE_CONFIG).write(path)
        wb = load_workbook(path)
        ws = wb.active
        # 找数据行:第3列是箱数(数值 int)
        data_rows = [r for r in range(1, ws.max_row + 1)
                     if isinstance(ws.cell(r, 3).value, (int, float))
                     and ws.cell(r, 4).value is not None]
        assert data_rows, "未找到 PL 数据行"
        row_nets = []
        for r in data_rows:
            c, d, f_col = ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 6).value
            assert round(d * c, 2) == pytest.approx(f_col, abs=0.005), (
                f"行{r}: {d}×{c}={round(d*c,2)} != 净重 {f_col}")
            row_nets.append(f_col)
        # 合计行 = Σ行净重(合计行第3列也是数值,取比数据行更靠后的)
        total_row = max(data_rows) + 1
        total_net = ws.cell(total_row, 6).value
        assert round(sum(row_nets), 2) == pytest.approx(total_net, abs=0.01)
    finally:
        Path(path).unlink(missing_ok=True)


# ── B3:CI 净重优先读 PL 回写 ─────────────────────────────────────


def test_ci_reads_pl_derived_net_weight():
    # PL 已回写 derived.total_net_weight 时,CI 必须用它,而非独立累加
    model = make_resolved_model(with_prices=True, with_weights=True)
    model.derived.total_net_weight = 123.45  # PL 的自洽口径(≠独立累加的 570)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        CIWriter(model, SAMPLE_CONFIG).write(path)
        wb = load_workbook(path)
        ws = wb.active
        nw_text = None
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "N.W.:" in c.value:
                    nw_text = c.value
        assert nw_text and "123.45" in nw_text, f"CI 未用 PL 回写净重: {nw_text!r}"
    finally:
        Path(path).unlink(missing_ok=True)


def test_ci_falls_back_when_pl_not_run():
    # derived 为 None(PL 未运行)时回退独立累加,行为不回归
    model = make_resolved_model(with_prices=True, with_weights=True)
    assert model.derived.total_net_weight is None
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        info = CIWriter(model, SAMPLE_CONFIG).write(path)
        assert info["total_net_weight"] == pytest.approx(570.0)
    finally:
        Path(path).unlink(missing_ok=True)


def test_ci_pl_net_weight_same_source_end_to_end():
    """端到端:先 PL(回写 derived)再 CI,两单净重严格相等。"""
    model = make_resolved_model(with_prices=True, with_weights=True)
    with tempfile.TemporaryDirectory() as td:
        pl_path = str(Path(td) / "pl.xlsx")
        ci_path = str(Path(td) / "ci.xlsx")
        pl_info = PLWriterLite(model, SAMPLE_CONFIG).write(pl_path)
        ci_info = CIWriter(model, SAMPLE_CONFIG).write(ci_path)
        assert ci_info["total_net_weight"] == pytest.approx(pl_info["total_net_weight"])
        assert check_ci_pl_net_weight(ci_info, pl_info) is None


# ── B3:跨单净重校验函数 ──────────────────────────────────────────


def test_net_weight_check_mismatch_warns():
    warn = check_ci_pl_net_weight(
        {"total_net_weight": 570.0}, {"success": True, "total_net_weight": 560.0})
    assert warn and "净重" in warn


def test_net_weight_check_skips_when_pl_failed():
    assert check_ci_pl_net_weight({"total_net_weight": 570.0}, {"success": False}) is None
    assert check_ci_pl_net_weight(None, None) is None
    assert check_ci_pl_net_weight(
        {"total_net_weight": 570.0}, {"success": True}) is None  # PL 无净重字段
