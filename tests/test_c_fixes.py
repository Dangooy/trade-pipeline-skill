"""C 系列(第三批审计遗留修复)的回归测试。

C1 LLM 截断标记 / C2 assembler 不丢弃降级标记 / C3 CI 吨数量格式 /
C4 欧洲数字格式解析。
"""
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from tests.conftest import make_resolved_model, SAMPLE_CONFIG
from trade_pipeline.understanding.llm_parser import _to_float
from trade_pipeline.understanding.assembler import assemble
from trade_pipeline.writers.ci_writer import CIWriter
import pytest


# ── C4:欧洲数字格式 ──────────────────────────────────────────────


def test_european_full_format():
    assert _to_float("1.234,56") == pytest.approx(1234.56)


def test_european_decimal_comma_only():
    # "12,5" 只可能是小数逗号(美式千分位逗号后必是三位),此前被解析成 125
    assert _to_float("12,5") == pytest.approx(12.5)


def test_european_million_grouping():
    assert _to_float("1.234.567,89") == pytest.approx(1234567.89)


def test_us_thousands_still_works():
    assert _to_float("1,000") == 1000.0
    assert _to_float("1,000,000") == 1000000.0


def test_plain_float_unchanged():
    assert _to_float("1234.56") == pytest.approx(1234.56)
    assert _to_float("1000") == 1000.0


def test_negative_european():
    assert _to_float("-1.234,56") == pytest.approx(-1234.56)


def test_unparseable_returns_zero_and_warns(capsys):
    assert _to_float("N/A") == 0.0
    assert "WARNING" in capsys.readouterr().out  # 不再静默归零


def test_empty_and_none_strings():
    assert _to_float("") == 0.0
    assert _to_float(None) == 0.0
    assert _to_float("None") == 0.0


# ── C2:assembler 消费降级/截断标记 ────────────────────────────────


def _minimal_rfq(**extra):
    rfq = {
        "buyer_name_en": "Global Fasteners LLC",
        "currency": "USD",
        "price_unit": "USD/PC",
        "format": "standard",
        "items": [{"description": "BOLT M8", "quantity": 1000, "unit": "pcs"}],
        "source_file": "x.xlsx",
    }
    rfq.update(extra)
    return rfq


def _assemble(rfq, config):
    return assemble(rfq=rfq, config=config, order_no="T1",
                    hint_buyer_id="global_fasteners")


def test_clean_parse_stays_clean():
    from tests.conftest import SAMPLE_CONFIG
    model = _assemble(_minimal_rfq(), SAMPLE_CONFIG)
    assert model.meta.review_status == "clean"
    assert model.meta.parser_model == "rules"  # 无 _parser_model 标记 → 规则模式
    assert model.meta.parser_notes == ""


def test_degraded_flag_surfaces_in_meta():
    from tests.conftest import SAMPLE_CONFIG
    rfq = _minimal_rfq(_llm_degraded=True,
                       _llm_degraded_reason="LLM 响应解析失败: bad json")
    model = _assemble(rfq, SAMPLE_CONFIG)
    assert model.meta.review_status == "degraded"
    assert model.meta.parser_model == "llm_fallback_rules"
    assert "bad json" in model.meta.parser_notes


def test_truncated_flag_surfaces_in_meta():
    from tests.conftest import SAMPLE_CONFIG
    rfq = _minimal_rfq(_parser_model="llm", _llm_truncated=True)
    model = _assemble(rfq, SAMPLE_CONFIG)
    assert model.meta.review_status == "degraded"
    assert model.meta.parser_model == "llm"
    assert "4000" in model.meta.parser_notes


def test_llm_success_model_recorded():
    from tests.conftest import SAMPLE_CONFIG
    model = _assemble(_minimal_rfq(_parser_model="llm"), SAMPLE_CONFIG)
    assert model.meta.parser_model == "llm"
    assert model.meta.review_status == "clean"  # 成功且未截断不算降级


# ── C3:CI 吨数量两位小数 ─────────────────────────────────────────


def _find_qty_header_row(ws):
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 14).value or "") in ("QTY\n数量", "QTY"):
            return r
    raise AssertionError("qty header row not found")


def test_ci_ton_qty_two_decimals():
    model = make_resolved_model(with_prices=True)
    model.order.format = "washers_mar"
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        CIWriter(model, SAMPLE_CONFIG).write(path)
        wb = load_workbook(path)
        ws = wb.active
        # 找一个数量数据行:第14列数值且第3列有描述
        for r in range(1, ws.max_row + 1):
            if isinstance(ws.cell(r, 14).value, (int, float)) and ws.cell(r, 3).value:
                fmt = ws.cell(r, 14).number_format
                assert fmt == "#,##0.00", f"吨计价数量格式仍为 {fmt}"
                return
        raise AssertionError("no qty data row found")
    finally:
        Path(path).unlink(missing_ok=True)


def test_ci_pcs_qty_stays_integer():
    model = make_resolved_model(with_prices=True)
    assert model.order.format == "standard"
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        CIWriter(model, SAMPLE_CONFIG).write(path)
        wb = load_workbook(path)
        ws = wb.active
        for r in range(1, ws.max_row + 1):
            if isinstance(ws.cell(r, 14).value, (int, float)) and ws.cell(r, 3).value:
                assert ws.cell(r, 14).number_format == "#,##0"
                return
        raise AssertionError("no qty data row found")
    finally:
        Path(path).unlink(missing_ok=True)


def test_ci_per_ton_price_unit_also_two_decimals():
    # format=standard 但 price_unit 含 TON:吨计价同样需要两位小数
    model = make_resolved_model(with_prices=True)
    model.order.price_unit = "USD/TON"
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name
    try:
        CIWriter(model, SAMPLE_CONFIG).write(path)
        wb = load_workbook(path)
        ws = wb.active
        for r in range(1, ws.max_row + 1):
            if isinstance(ws.cell(r, 14).value, (int, float)) and ws.cell(r, 3).value:
                assert ws.cell(r, 14).number_format == "#,##0.00"
                return
    finally:
        Path(path).unlink(missing_ok=True)
